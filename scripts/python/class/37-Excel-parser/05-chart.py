#!/usr/bin/python

from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.chart import (
    Reference,
    Series,
    BarChart
)

wb = load_workbook('shared/revenue.xlsx')
wsheet = wb.get_sheet_by_name('sales')

#data =   Reference(wsheet, min_col=5, min_row=2, max_col=5, max_row=10)
#categs = Reference(wsheet, min_col=3, min_row=2, max_col=3, max_row=10)

data =   Reference(wsheet, (5, 2), (5, 10))
categs = Reference(wsheet, (3, 2), (3, 10))

chart = BarChart()
chart.add_data(data=data)
chart.set_categories(categs)

chart.legend = None
chart.y_axis.majorGridlines = None
chart.varyColors = True
chart.title = "Sales By name"

wsheet.add_chart(chart, "H2")    

wb.save("shared/revenue.xlsx")
