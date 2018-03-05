#!/usr/bin/python

from openpyxl import load_workbook
from openpyxl import Workbook

from openpyxl.chart import (
    Reference,
    BarChart,
    ProjectedPieChart
)

wb = load_workbook('shared/revenue.xlsx')
wsheet = wb.get_sheet_by_name('population')

data =   Reference(wsheet, min_col=2, min_row=2, max_col=2, max_row=12)
categary = Reference(wsheet, min_col=1, min_row=2, max_col=1, max_row=12)

chart = BarChart()
chart.add_data(data=data)
chart.set_categories(categary)

chart.legend = None
chart.y_axis.majorGridlines = None
#chart.varyColors = True
chart.title = "World Population graph"
wsheet.add_chart(chart, "H2")    

projected_pie = ProjectedPieChart()
projected_pie.type = "pie"
projected_pie.splitType = "val" # split by value

data =   Reference(wsheet, min_col=2, min_row=2, max_col=2, max_row=12)
labels = Reference(wsheet, min_col=1, min_row=2, max_col=1, max_row=12)

projected_pie.add_data(data, titles_from_data=True)
projected_pie.title = "Population graph"
projected_pie.set_categories(labels)
wsheet.add_chart(projected_pie, "H15")
wb.save("shared/revenue.xlsx")
